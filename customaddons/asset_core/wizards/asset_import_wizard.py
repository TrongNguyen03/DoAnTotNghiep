# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError
import base64
import xlrd
import openpyxl
from datetime import datetime, date

class AssetImportWizard(models.TransientModel):
    _name = "asset.import.wizard"
    _description = "Import Tài sản từ Excel (hỗ trợ đầy đủ các trường của asset.asset)"

    file = fields.Binary(string="File Excel (.xls/.xlsx)", required=True)
    filename = fields.Char(string="Tên file")

    def _parse_date(self, cell, workbook=None):
        """Chuyển cell Excel sang datetime.date; hỗ trợ số Excel và chuỗi định dạng phổ biến."""
        if cell is None or cell == '':
            return False
        # Excel serial khi đọc bằng xlrd
        if workbook and isinstance(cell, (float, int)):
            try:
                return datetime(*xlrd.xldate_as_tuple(cell, workbook.datemode)).date()
            except Exception:
                return False

        if isinstance(cell, str):
            s = cell.strip()
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
                try:
                    return datetime.strptime(s, fmt).date()
                except Exception:
                    pass
            try:
                return date.fromisoformat(s)
            except Exception:
                return False
        # openpyxl có thể trả datetime
        if isinstance(cell, datetime):
            return cell.date()
        return False

    def _parse_float(self, cell):
        if cell is None or cell == '':
            return 0.0
        if isinstance(cell, (int, float)):
            return float(cell)
        s = str(cell).strip().replace(',', '')
        try:
            return float(s)
        except Exception:
            return 0.0

    def _parse_string(self, cell):
        if cell is None:
            return ''
        if isinstance(cell, float):
            if cell.is_integer():
                return str(int(cell))
            return str(cell)
        return str(cell).strip()

    def action_import(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_("Vui lòng chọn file Excel để import."))

        data = base64.b64decode(self.file)
        workbook = None
        sheet = None
        is_xlsx = False

        # Thử đọc bằng openpyxl (xlsx)
        try:
            wb = openpyxl.load_workbook(filename=None, file_contents=data, data_only=True)
            sheet = wb.active
            is_xlsx = True
        except Exception:
            # Nếu không được thì fallback sang xlrd (xls)
            try:
                workbook = xlrd.open_workbook(file_contents=data)
                sheet = workbook.sheet_by_index(0)
                is_xlsx = False
            except Exception as e:
                raise UserError(_("Không thể đọc file Excel: %s") % e)

        Asset = self.env['asset.asset']
        Category = self.env['asset.category']
        Employee = self.env['hr.employee']
        Department = self.env['hr.department']

        warnings = []
        errors = []
        created = 0

        state_map = {
            'stock': 'stock', 'trong kho': 'stock',
            'in_use': 'in_use', 'đang sử dụng': 'in_use', 'sử dụng': 'in_use',
            'maintenance': 'maintenance', 'đang bảo trì': 'maintenance',
            'broken': 'broken', 'hỏng': 'broken', 'scrap': 'broken'
        }

        # Số cột tối thiểu
        def get_ncols():
            return sheet.max_column if is_xlsx else sheet.ncols

        def get_nrows():
            return sheet.max_row if is_xlsx else sheet.nrows

        if get_ncols() < 11:
            raise UserError(_("File có ít hơn 11 cột theo template. Vui lòng dùng template chuẩn!"))

        for rowx in range(2, get_nrows() + 1) if is_xlsx else range(1, sheet.nrows):
            try:
                def _cell(col):
                    try:
                        if is_xlsx:
                            return sheet.cell(row=rowx, column=col + 1).value
                        else:
                            return sheet.cell_value(rowx, col)
                    except Exception:
                        return ''

                name = _cell(0) or ''
                code = self._parse_string(_cell(1)) or ''
                serial = self._parse_string(_cell(2)) or ''
                category_name = _cell(3) or ''
                employee_name = _cell(4) or ''
                department_name = _cell(5) or ''
                purchase_date_cell = _cell(6)
                purchase_value_cell = _cell(7)
                warranty_cell = _cell(8)
                state_raw = _cell(9) or ''
                note = _cell(10) or ''

                if not name:
                    errors.append(_("Dòng %s: Thiếu Tên tài sản, bỏ qua.") % rowx)
                    continue

                # Kiểm tra trùng code/serial
                if code and Asset.search([('code', '=', code)], limit=1):
                    warnings.append(_("Dòng %s: Mã tài sản '%s' đã tồn tại — bỏ qua.") % (rowx, code))
                    continue
                if serial and Asset.search([('serial_number', '=', serial)], limit=1):
                    warnings.append(_("Dòng %s: Số seri '%s' đã tồn tại — bỏ qua.") % (rowx, serial))
                    continue

                # Category
                category_id = False
                if category_name:
                    cat = Category.search([('name', '=', str(category_name).strip())], limit=1)
                    if not cat:
                        cat = Category.create({'name': str(category_name).strip()})
                    category_id = cat.id

                # Employee
                employee_id = False
                emp_rec = False
                if employee_name:
                    emp = Employee.search([('name', '=', str(employee_name).strip())], limit=1)
                    if not emp:
                        emp = Employee.search([('work_email', '=', str(employee_name).strip())], limit=1)
                    if emp:
                        employee_id = emp.id
                        emp_rec = emp
                    else:
                        warnings.append(_("Dòng %s: Không tìm thấy nhân viên '%s' — để trống.") % (rowx, employee_name))

                # Department
                department_id = False
                if department_name:
                    dept = Department.search([('name', '=', str(department_name).strip())], limit=1)
                    if dept:
                        department_id = dept.id
                    else:
                        warnings.append(_("Dòng %s: Không tìm thấy phòng ban '%s' — để trống.") % (rowx, department_name))
                elif emp_rec and emp_rec.department_id:
                    department_id = emp_rec.department_id.id

                # Dates + value
                purchase_date = self._parse_date(purchase_date_cell, workbook if not is_xlsx else None)
                warranty_end_date = self._parse_date(warranty_cell, workbook if not is_xlsx else None)
                purchase_value = self._parse_float(purchase_value_cell)

                # State
                state_val = 'stock'
                if state_raw:
                    k = str(state_raw).strip().lower()
                    state_val = state_map.get(k, 'stock')

                vals = {
                    'name': str(name).strip(),
                    'code': code or False,
                    'serial_number': serial or False,
                    'category_id': category_id or False,
                    'employee_id': employee_id or False,
                    'department_id': department_id or False,
                    'purchase_date': purchase_date or False,
                    'purchase_value': purchase_value or 0.0,
                    'warranty_end_date': warranty_end_date or False,
                    'state': state_val,
                    'note': str(note).strip() if note else False,
                }

                try:
                    Asset.create(vals)
                    created += 1
                except Exception as e:
                    errors.append(_("Dòng %s: Lỗi khi tạo tài sản: %s") % (rowx, e))

            except Exception as e:
                errors.append(_("Dòng %s: Lỗi khi đọc: %s") % (rowx, e))
                continue

        msgs = []
        msgs.append(_("Số tài sản tạo thành công: %s") % created)
        if warnings:
            msgs.append(_("Cảnh báo:"))
            msgs += warnings
        if errors:
            msgs.append(_("Lỗi:"))
            msgs += errors

        message = "\n".join(msgs[:200])

        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("Kết quả import"),
                "message": message,
                "sticky": bool(errors or warnings),
            },
        }
